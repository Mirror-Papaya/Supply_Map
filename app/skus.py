"""SKU 字典：SKU、品名、别名（旧码/平台码，用于迁移与识别纠错）。"""
import csv
import io
import re

from . import db


def key(s: str) -> str:
    return (s or "").strip().upper()


def lookup(sku: str) -> dict | None:
    """精确查找（不区分大小写），也匹配别名。返回字典行。"""
    k = key(sku)
    if not k:
        return None
    row = db.query_one("SELECT * FROM sku WHERE UPPER(sku)=?", (k,))
    if row:
        return row
    for r in db.query("SELECT * FROM sku WHERE aliases LIKE ?", (f"%{sku.strip()}%",)):
        if k in {key(a) for a in split_aliases(r["aliases"])}:
            return r
    return None


def split_aliases(s: str) -> list:
    return [a for a in re.split(r"[,，;；|\s]+", s or "") if a]


def names(skus) -> dict:
    skus = list(skus)
    if not skus:
        return {}
    marks = ",".join("?" * len(skus))
    return {r["sku"]: r["name"] for r in db.query(f"SELECT sku, name FROM sku WHERE sku IN ({marks})", skus)}


def count() -> int:
    return db.query_one("SELECT COUNT(*) n FROM sku")["n"]


def upsert(rows: list) -> int:
    """rows: [{"sku":..., "name":..., "aliases":...}]；已存在的 SKU 更新品名，别名合并。"""
    n = 0
    now = db.now()
    with db.tx() as c:
        for r in rows:
            sku = (r.get("sku") or "").strip()
            if not sku:
                continue
            name = (r.get("name") or "").strip()
            aliases = split_aliases(r.get("aliases") or "")
            old = c.execute("SELECT name, aliases FROM sku WHERE sku=?", (sku,)).fetchone()
            if old:
                merged = split_aliases(old["aliases"])
                merged += [a for a in aliases if a not in merged]
                c.execute("UPDATE sku SET name=?, aliases=?, updated_at=? WHERE sku=?",
                          (name or old["name"], ",".join(merged), now, sku))
            else:
                c.execute("INSERT INTO sku(sku, name, aliases, updated_at) VALUES(?,?,?,?)",
                          (sku, name, ",".join(aliases), now))
            n += 1
    return n


# ---------- 表格文件读取（SKU 字典导入 / Excel 拣货单共用） ----------

SKU_HEADERS = ["商家sku", "卖家sku", "seller sku", "merchant sku", "商家编码", "sku编码", "sku", "货号", "商品编码",
               "编码", "item code", "item no", "code", "product code", "model"]
NAME_HEADERS = ["商品名称", "品名", "名称", "产品名称", "product name", "item name", "name", "description", "描述"]
QTY_HEADERS = ["拣货数量", "数量", "qty", "quantity", "件数", "pcs"]
ALIAS_HEADERS = ["别名", "旧码", "旧sku", "原码", "alias", "aliases", "old sku"]


def read_table(filename: str, data: bytes) -> list:
    """读取 xlsx / xls / csv 的第一个工作表，返回二维字符串列表。"""
    name = filename.lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        for enc in ("utf-8-sig", "gbk", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]
    if name.endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        sh = book.sheet_by_index(0)
        return [[_cell_str(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return [[_cell_str(v) for v in row] for row in ws.iter_rows(values_only=True)]


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _match_header(cell: str, candidates: list) -> int:
    c = re.sub(r"\s+", " ", cell.strip().lower())
    for i, h in enumerate(candidates):
        if c == h:
            return i
    for i, h in enumerate(candidates):
        if h in c:
            return i + len(candidates)
    return -1


def detect_columns(table: list, want=("sku", "name", "qty", "aliases"), max_scan=15) -> tuple:
    """在前几行里找表头，返回 (表头行号, {字段: 列号})。"""
    spec = {"sku": SKU_HEADERS, "name": NAME_HEADERS, "qty": QTY_HEADERS, "aliases": ALIAS_HEADERS}
    best = (-1, {})
    for r, row in enumerate(table[:max_scan]):
        found = {}
        for field in want:
            scored = [(s, ci) for ci, cell in enumerate(row) if cell
                      for s in [_match_header(cell, spec[field])] if s >= 0]
            used = set(found.values())
            scored = [x for x in scored if x[1] not in used]
            if scored:
                found[field] = min(scored)[1]
        if "sku" in found and len(found) > len(best[1]):
            best = (r, found)
    return best


def rows_from_table(table: list, want=("sku", "name", "qty", "aliases")) -> list:
    hr, cols = detect_columns(table, want)
    if hr < 0:
        return []
    out = []
    for row in table[hr + 1:]:
        get = lambda f: row[cols[f]].strip() if f in cols and cols[f] < len(row) else ""
        sku = get("sku")
        if not sku or sku.lower() in ("合计", "total", "总计"):
            continue
        out.append({"sku": sku, "name": get("name"), "qty": get("qty"), "aliases": get("aliases")})
    return out
