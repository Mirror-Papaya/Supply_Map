"""查询结果 → 飞书消息卡片（JSON 2.0）/ 纯文本兜底 / Excel 附件。"""
import io

MAX_TABLE_ROWS = 80
FLAG_COLOR = {"余量不足": "red", "未登记": "red", "多库位": "orange", "识别纠错": "blue"}


def _qty_text(row) -> str:
    if row["qty"] is None:
        return "—"
    when = (row["qty_at"] or "")[5:10]
    return f"{row['qty']}（{when} {row['qty_by'] or ''}）".replace(" ）", "）")


def _need(it) -> str:
    return "" if it["qty"] is None else str(it["qty"])


def table_rows(res) -> list:
    rows = []
    for ln in res["lines"]:
        it, r = ln["item"], ln["row"]
        rows.append({"seq": str(ln["seq"]), "loc": r["loc_code"], "sku": it["sku"], "name": it["name"] or "",
                     "need": _need(it), "qty": _qty_text(r),
                     "flags": [{"text": f, "color": FLAG_COLOR.get(f, "neutral")} for f in it["flags"]]})
    for it in res["unregistered"]:
        rows.append({"seq": "—", "loc": "未登记", "sku": it["sku"], "name": it["name"] or "", "need": _need(it),
                     "qty": "—", "flags": [{"text": f, "color": FLAG_COLOR.get(f, "neutral")} for f in it["flags"]]})
    return rows


def summary_lines(res) -> list:
    shelves = sum(len(v) for v in res["warehouses"].values())
    out = [f"识别 {res['demand_rows']} 行 → {len(res['items'])} 个 SKU，分布在 {shelves} 个货架（{len(res['lines'])} 个库位）"]
    short = [it for it in res["items"] if "余量不足" in it["flags"]]
    if short:
        out.append("⚠ 余量不足（按最近盘点数，仅供参考）：" + "、".join(f"{it['sku']}（需 {it['qty']} / 余 {it['total']}）"
                                                         for it in short))
    if res["unregistered"]:
        out.append("⚠ 未登记库位：" + "、".join(it["sku"] for it in res["unregistered"]))
    multi = [it for it in res["items"] if "多库位" in it["flags"]]
    if multi:
        out.append("ℹ 多库位：" + "、".join(f"{it['sku']}（{len(it['locations'])} 处）" for it in multi))
    if res["unmatched"]:
        parts = []
        for u in res["unmatched"]:
            s = u["input"] or u["name"] or "?"
            if u["candidates"]:
                s += f"（可能是 {'/'.join(u['candidates'][:3])}）"
            parts.append(s)
        out.append("❓ 没找到的 SKU：" + "、".join(parts))
    fuzzy = [it for it in res["items"] if it["fuzzy"]]
    if fuzzy:
        out.append("🔎 已自动纠错：" + "、".join(f"{'/'.join(sorted(set(it['inputs'])))} → {it['sku']}" for it in fuzzy))
    return out


def result_card(res, image_keys: dict, source: str = "") -> dict:
    problems = bool(res["unregistered"] or res["unmatched"] or any("余量不足" in it["flags"] for it in res["items"]))
    rows = table_rows(res)
    elements = [{"tag": "markdown", "content": "\n".join(summary_lines(res))}]
    if rows:
        elements.append({
            "tag": "table",
            "page_size": 10,
            "row_height": "low",
            "freeze_first_column": True,
            "header_style": {"text_align": "left", "text_size": "normal", "background_style": "grey",
                             "bold": True, "lines": 1},
            "columns": [
                {"name": "seq", "display_name": "顺序", "data_type": "text", "width": "56px"},
                {"name": "loc", "display_name": "库位", "data_type": "text", "width": "auto"},
                {"name": "sku", "display_name": "SKU", "data_type": "text", "width": "auto"},
                {"name": "name", "display_name": "品名", "data_type": "text", "width": "auto"},
                {"name": "need", "display_name": "需求", "data_type": "text", "width": "56px"},
                {"name": "qty", "display_name": "余量（盘点）", "data_type": "text", "width": "auto"},
                {"name": "flags", "display_name": "提醒", "data_type": "options", "width": "auto"},
            ],
            "rows": rows[:MAX_TABLE_ROWS],
        })
        if len(rows) > MAX_TABLE_ROWS:
            elements.append({"tag": "markdown", "content": f"表格只显示前 {MAX_TABLE_ROWS} 行，完整结果见附件 Excel。"})
    for wh, key in sorted(image_keys.items()):
        elements.append({"tag": "markdown", "content": f"**仓库 {wh} 定位图**（红色为本次货架，数字为拣货顺序）"})
        elements.append({"tag": "img", "img_key": key, "alt": {"tag": "plain_text", "content": f"仓库 {wh} 定位图"},
                         "scale_type": "fit_horizontal", "preview": True})
    elements.append({"tag": "markdown", "content": "<font color='grey'>余量为最近一次盘点数（日期 盘点人），仅供参考</font>"})
    title = f"拣货定位 · {len(res['items'])} 个 SKU" + (f"（{source}）" if source else "")
    return {
        "schema": "2.0",
        "config": {"width_mode": "fill"},
        "header": {"title": {"tag": "plain_text", "content": title}, "template": "orange" if problems else "blue"},
        "body": {"elements": elements},
    }


def result_text(res) -> str:
    """卡片发送失败时的纯文本兜底。"""
    lines = summary_lines(res) + [""]
    for r in table_rows(res):
        flags = " ".join(f"[{f['text']}]" for f in r["flags"])
        lines.append(f"{r['seq']}. {r['loc']}  {r['sku']} {r['name'][:16]}  需{r['need'] or '?'}  余{r['qty']} {flags}")
    return "\n".join(lines)


def result_xlsx(res) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "拣货定位"
    head = ["顺序", "库位", "SKU", "品名", "需求", "余量(盘点)", "盘点时间", "盘点人", "提醒"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E9ECEF")
    for ln in res["lines"]:
        it, r = ln["item"], ln["row"]
        ws.append([ln["seq"], r["loc_code"], it["sku"], it["name"], it["qty"], r["qty"], r["qty_at"], r["qty_by"],
                   " ".join(it["flags"])])
    for it in res["unregistered"]:
        ws.append(["", "未登记", it["sku"], it["name"], it["qty"], None, None, None, " ".join(it["flags"])])
    for u in res["unmatched"]:
        ws.append(["", "没找到SKU", u["input"], u["name"], u["qty"], None, None, None,
                   ("可能是 " + "/".join(u["candidates"])) if u["candidates"] else ""])
    for col, w in zip("ABCDEFGHI", (6, 14, 18, 30, 6, 10, 18, 10, 20)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
